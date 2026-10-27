#!/usr/bin/env python3
"""
Converts a POC label sheet between CSV (what label-sheets.ts reads) and a friendlier XLSX with drop-downs, so a labeller cannot mistype a value.

  python scripts/poc-event-insight/sheet_xlsx.py to-xlsx local/label-sheet-owner-tuning-sample.csv local/label-sheet-owner-tuning-sample.xlsx
  python scripts/poc-event-insight/sheet_xlsx.py to-csv  local/label-sheet-owner-tuning-sample.xlsx local/label-sheet-owner-tuning-sample.csv

Needs openpyxl (pip install openpyxl). The allowed values are the ones in label-sheets-lib.ts and labelling-guide-2026-10-21.md; keep them in step.
The sheets hold third-party captions: keep the files in the git-ignored local/ folder.
"""
import csv
import math
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

CONTEXT = ['id', 'account', 'stratum', 'url', 'post_date', 'caption', 'hashtags', 'tagged_place']
LABELS = ['is_event', 'audience_country', 'geo_scope', 'post_language', 'target_area_level', 'target_area_name', 'target_area_source',
          'eligibility', 'track_count', 'age_written', 'cost_tier', 'cost_source', 'sensitive', 'unsure', 'notes']
COLUMNS = CONTEXT + LABELS

COUNTRIES = ['id', 'my', 'sg', 'au', 'dk', 'th', 'jp', 'kr', 'cn', 'in', 'ph', 'vn', 'us', 'gb', 'eg', 'international', 'unclear']
LISTS = {
    'is_event': ['yes', 'no'],
    'cost_source': ['caption', 'image', 'none'],
    'audience_country': COUNTRIES,
    'geo_scope': ['venue_local', 'city', 'region', 'national', 'international', 'online', 'unspecified', 'unclear'],
    'post_language': ['id', 'en', 'ms', 'mixed', 'other', 'unclear'],
    'target_area_level': ['neighborhood', 'city', 'regency', 'province', 'country', 'continent', 'none', 'unclear'],
    'target_area_source': ['stated', 'stated_poster', 'none', 'unclear'],
    'eligibility': ['open_to_public', 'restricted', 'unspecified', 'unclear'],
    'age_written': ['yes', 'no', 'unclear'],
    'cost_tier': ['free', 'low', 'medium', 'high', 'unspecified', 'unclear'],
    'sensitive': ['none', 'religion', 'politics_civic', 'health_condition', 'sexuality', 'ethnicity', 'unclear'],
    'unsure': ['yes', 'no'],
}

GUIDE = [
    ('Field', 'What to write'),
    ('is_event', 'yes if the post advertises or reports on something with a date (a schedule): an event that is coming up or running. no if it has no schedule: a recap or thank-you after it ended, a repost of a personal vlog, news, a shop or restaurant promo, a quote or greeting. For a "no" post fill only audience_country, post_language, sensitive and unsure; write unclear in the other label columns and say why in notes. If it is about a known event (a recap or update), say which in notes.'),
    ('audience_country', 'The country the post is written FOR: two-letter code (id, my, sg, au, dk ...), "international", or "unclear". Steps: (1) does the post state who it is for? (2) else language + where the account is based; an Indonesian post from an Indonesian account is "id" even if the event is abroad; (3) shared language (English, Malay, Indonesian) needs a second clue: currency, phone code, time zone, account; (4) online with no stated audience = international; (5) last resort: where the event takes place.'),
    ('geo_scope', 'venue_local = a venue is named (the audience is the people around it). city = only a city or area is named, no venue. region = a province, or several named areas. national = invites people from anywhere in the country ("seluruh Indonesia"). international. online = remote and open to all. unspecified = no basis. "Open to the public" is NOT national. A venue called "Galeri Nasional" is just a venue.'),
    ('post_language', 'id, en, ms (Malay), mixed (Indonesian and English in one caption), other.'),
    ('target_area_level', 'The level of the area you wrote in target_area_name: neighborhood (ONLY when the post says RW, RT, kampung, dusun or lingkungan), city, regency, province, country, continent. If the name could be a city or its province (Jogja, Yogyakarta, Jakarta), pick the WIDER one. none when target_area_source is none.'),
    ('target_area_name', 'The area the post SAYS the event is for, as written: "Jogja", "DIY", "RW 05 Kelurahan X". Read the caption AND the poster. An address or "Tempat:" line, the venue name, and where performers, exhibitors or vendors come from are NOT an audience. People who attend, register or compete ARE ("peserta se-DIY", "warga Jogja"). Several areas: separate them with ; . Write none if the post states no area. Do NOT write an area taken from the venue, the hashtags or the account: the system derives those.'),
    ('target_area_source', 'stated = the caption says it. stated_poster = only the poster says it. none = nothing states an area (the system will derive it). Do not label venue, hashtag or account areas by hand.'),
    ('eligibility', 'restricted ONLY if membership limits who may take part (invitation only, one school, one company, members only). Otherwise open_to_public: an event with no restriction is open. unspecified only for a post that is not an event. Age and gender are NOT eligibility.'),
    ('track_count', 'How many separately qualifying categories the event lists. If it lists none, it is one tier open to everybody: write 1. "Kids Dash 6-10, 5K 12+, 10K 17+" is 3. Ticket packs and prices are not categories. For a post that is not an event write unclear.'),
    ('age_written', 'yes if any age NUMBER is written next to an age word (usia 6-10 tahun, 12+, under 17). Do not infer from SD, SMA or mahasiswa.'),
    ('cost_tier', 'free, low (under about Rp 50,000), medium (up to about Rp 250,000), high (above), unspecified. The cheapest fee a participant pays. You may read the price from the image.'),
    ('cost_source', 'Where you read the price: caption, image (the price is only in the picture), or none (no price anywhere). The model sees the caption only, so cost is scored on caption and none rows.'),
    ('sensitive', 'none, or the topic the event is about: religion, politics_civic, health_condition, sexuality, ethnicity.'),
    ('unsure', 'yes or no. Write unclear in a field you cannot decide after about two minutes, and say why in notes. Never guess.'),
    ('Rules', 'Label from the post only (caption, hashtags, tagged place, text on the image). Work alone; do not look at any other labels or model output. The posts are content of other people: do not copy or share them.'),
]


def read_csv(path):
    with open(path, encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))


def to_xlsx(src, dst):
    rows = read_csv(src)
    wb = Workbook()
    ws = wb.active
    ws.title = 'Labels'

    head_fill = PatternFill('solid', fgColor='DDE7F3')
    label_fill = PatternFill('solid', fgColor='FFF7D6')
    for c, name in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=name)
        cell.font = Font(bold=True)
        cell.fill = head_fill if name in CONTEXT else PatternFill('solid', fgColor='F5D98A')
        cell.alignment = Alignment(vertical='center', wrap_text=True)

    widths = {'id': 12, 'account': 18, 'stratum': 12, 'url': 14, 'post_date': 12, 'caption': 90, 'hashtags': 28, 'tagged_place': 22,
              'is_event': 9, 'cost_source': 11, 'audience_country': 14, 'geo_scope': 14, 'post_language': 12, 'target_area_level': 16, 'target_area_name': 26, 'target_area_source': 16, 'eligibility': 16,
              'track_count': 11, 'age_written': 11, 'cost_tier': 12, 'sensitive': 16, 'unsure': 9, 'notes': 36}
    for c, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(c)].width = widths[name]

    for r, row in enumerate(rows, start=2):
        for c, name in enumerate(COLUMNS, start=1):
            val = (row.get(name) or '')
            if name == 'post_date':
                val = val[:10]
            cell = ws.cell(row=r, column=c, value=val)
            cell.alignment = Alignment(vertical='top', wrap_text=name in ('caption', 'hashtags', 'tagged_place', 'notes', 'account'))
            if name in LABELS:
                cell.protection = Protection(locked=False)
                cell.fill = label_fill
        ws.row_dimensions[r].height = min(300, max(45, 15 * math.ceil(len(row.get('caption') or '') / 95)))

    last = len(rows) + 1
    col = {name: get_column_letter(i) for i, name in enumerate(COLUMNS, start=1)}

    for name, values in LISTS.items():
        dv = DataValidation(type='list', formula1='"' + ','.join(values) + '"', allow_blank=True)
        dv.showErrorMessage = True
        dv.errorTitle = 'Not an allowed value'
        dv.error = 'Pick one from the list (see the Guide sheet).'
        if name == 'audience_country':
            dv.errorStyle = 'warning'  # a country not on the short list may be typed; label-sheets.ts check accepts any two-letter code
            dv.error = 'Not on the short list. Continue only if it is a correct two-letter country code (lower case).'
        dv.promptTitle = name
        dv.prompt = 'Choose from the list'
        dv.showInputMessage = False
        ws.add_data_validation(dv)
        dv.add(f'{col[name]}2:{col[name]}{last}')

    tc = col['track_count']
    dv = DataValidation(type='custom',
                        formula1=f'=OR(AND(ISNUMBER({tc}2),{tc}2>=0,{tc}2<=30,{tc}2=INT({tc}2)),{tc}2="unclear")', allow_blank=True)
    dv.showErrorMessage = True
    dv.errorTitle = 'track_count'
    dv.error = 'A whole number from 0 to 30, or the word unclear.'
    ws.add_data_validation(dv)
    dv.add(f'{tc}2:{tc}{last}')

    # a still-empty label cell shows in a stronger yellow so nothing is skipped
    first_label, last_label = col['is_event'], col['unsure']
    ws.conditional_formatting.add(
        f'{first_label}2:{last_label}{last}',
        FormulaRule(formula=[f'LEN({first_label}2)=0'], fill=PatternFill('solid', bgColor='FFE08A', fgColor='FFE08A')))

    ws.freeze_panes = 'B2'
    ws.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{last}'
    ws.protection.sheet = True  # no password: stops an accidental edit of the caption; label cells are unlocked
    ws.protection.formatColumns = False
    ws.protection.formatRows = False
    ws.protection.autoFilter = False
    ws.protection.sort = False

    g = wb.create_sheet('Guide')
    for i, (a, b) in enumerate(GUIDE, start=1):
        g.cell(row=i, column=1, value=a).font = Font(bold=(i == 1 or a == 'Rules'))
        c = g.cell(row=i, column=2, value=b)
        c.alignment = Alignment(wrap_text=True, vertical='top')
        g.cell(row=i, column=1).alignment = Alignment(vertical='top')
    g.column_dimensions['A'].width = 20
    g.column_dimensions['B'].width = 120

    wb.save(dst)
    print(f'{len(rows)} rows written to {dst}')


def to_csv(src, dst):
    wb = load_workbook(src, data_only=True)
    ws = wb['Labels']
    header = [c.value for c in ws[1]]
    missing = [c for c in CONTEXT if c not in header]
    if missing:
        raise SystemExit(f'missing columns in {src}: {missing}')
    # by name, so a moved column still converts; a label column an older sheet lacks (is_event, cost_source) comes out blank
    idx = {name: header.index(name) for name in COLUMNS if name in header}
    out = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[idx['id']] in (None, ''):
            continue
        rec = {}
        for name in COLUMNS:
            v = row[idx[name]] if name in idx else None
            if v is None:
                v = ''
            elif isinstance(v, float) and v.is_integer():
                v = int(v)
            rec[name] = str(v).strip()
        out.append(rec)
    with open(dst, 'w', encoding='utf-8-sig', newline='') as f:
        w = csv.DictWriter(f, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(out)
    print(f'{len(out)} rows written to {dst}')


if __name__ == '__main__':
    if len(sys.argv) != 4 or sys.argv[1] not in ('to-xlsx', 'to-csv'):
        raise SystemExit(__doc__)
    {'to-xlsx': to_xlsx, 'to-csv': to_csv}[sys.argv[1]](sys.argv[2], sys.argv[3])
